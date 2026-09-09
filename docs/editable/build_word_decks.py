#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Editable Word decks reconstructed from USB PDFs:
  DOC-20260818-WA0010.pdf  VIP / Gold clinics pitch
  DOC-20260818-WA0011.pdf  Internal financial report
  DOC-20260818-WA0012.pdf  Small clinics / private practices pitch
  DOC-20260818-WA0013.pdf  Financial estimate + pricing strategy

Traffic figures that said ۱۰۰۰ visits/month are updated to ۵۰۰۰.
Prices, CAPEX/OPEX, and package fees are copied from the originals.
"""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x04, 0x1F, 0x45)
BLUE = RGBColor(0x12, 0x94, 0xE0)
INK = RGBColor(0x13, 0x23, 0x3C)
MUTED = RGBColor(0x5F, 0x6F, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xB4, 0x86, 0x1A)
ORANGE = RGBColor(0xC4, 0x56, 0x1A)
FONT = "Tahoma"
OUT = Path(__file__).resolve().parent


def _rpr_font(rPr, size_pt, bold=False, color=None, rtl=True):
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(size_pt * 2)))
    rPr.append(szCs)
    if bold:
        rPr.append(OxmlElement("w:b"))
        rPr.append(OxmlElement("w:bCs"))
    if color is not None:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "%02X%02X%02X" % (color[0], color[1], color[2]))
        rPr.append(c)
    if rtl:
        rPr.append(OxmlElement("w:rtl"))
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "fa-IR")
    lang.set(qn("w:bidi"), "fa-IR")
    rPr.append(lang)


def set_run(run, size=12, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    _rpr_font(rPr, size, bold=bold, color=tuple(color), rtl=True)


def set_para_rtl(p, align="right", space_after=8, space_before=0):
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.35
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    jc = OxmlElement("w:jc")
    val = "center" if align == "center" else ("left" if align == "left" else "right")
    jc.set(qn("w:val"), val)
    pPr.append(jc)


def add_text(doc, text, size=12, bold=False, color=INK, align="right", space_after=8, space_before=0):
    p = doc.add_paragraph()
    set_para_rtl(p, align=align, space_after=space_after, space_before=space_before)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="D0D7DE"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text, size=11, bold=False, color=INK, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    set_para_rtl(p, align="center" if center else "right", space_after=2, space_before=2)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def _table_bidi(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    vis = OxmlElement("w:bidiVisual")
    vis.set(qn("w:val"), "1")
    tblPr.append(vis)


def add_table(doc, headers, rows, header_fill="041F45", header_color=WHITE, zebra=("F8FBFC", "FFFFFF")):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _table_bidi(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, size=10, bold=True, color=header_color, fill=header_fill, center=True)
    for r_i, row in enumerate(rows):
        fill = zebra[0] if r_i % 2 == 0 else zebra[1]
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=10, color=INK, fill=fill, center=True)
    doc.add_paragraph()
    return table


def add_box(doc, lines, fill="FFF8E7", title=None, title_color=ORANGE, body_color=INK, size=11):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_bidi(table)
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill)
    set_cell_borders(cell, "E8C56B" if fill.upper().startswith("FFF") else "1A365D")
    cell.text = ""
    first = True
    if title:
        p = cell.paragraphs[0]
        set_para_rtl(p, align="right", space_after=6, space_before=4)
        run = p.add_run(title)
        set_run(run, size=12, bold=True, color=title_color)
        first = False
    for line in lines:
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        set_para_rtl(p, align="right", space_after=4, space_before=2)
        if isinstance(line, tuple):
            text, bold = line
        else:
            text, bold = line, False
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=body_color)
    doc.add_paragraph()
    return table


def new_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.right_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    section._sectPr.append(bidi)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    _rpr_font(style.element.get_or_add_rPr(), 12, rtl=True)
    return doc


def font_note(doc):
    add_text(
        doc,
        "فونت این فایل Tahoma است (روی ویندوز نصب است). برای تغییر فونت: Ctrl+A سپس وزیرمتن یا بی‌نازنین.",
        size=8,
        color=MUTED,
        align="center",
        space_after=0,
        space_before=6,
    )


def footer_ir(doc, extra=""):
    add_text(
        doc,
        "پلتفرم نوبت‌دهی سلام دکتر  |  www.salam-doctor.ir  |  salam-doctor.com"
        + (("  |  " + extra) if extra else ""),
        size=9,
        color=MUTED,
        align="center",
        space_before=10,
        space_after=2,
    )
    font_note(doc)


# ---------------------------------------------------------------------------
# WA0010 — VIP / Gold clinics
# ---------------------------------------------------------------------------
def build_vip():
    doc = new_doc()
    add_text(doc, "پروپوزال انحصاری توسعه بازار و جذب بیمار", size=20, bold=True, color=GOLD, align="center", space_after=4)
    add_text(
        doc,
        "ویژه مراکز درمانی و کلینیک‌های زیبایی رتبه یک — پلتفرم «سلام دکتر»",
        size=12,
        color=NAVY,
        align="center",
        space_after=14,
    )
    add_text(doc, "مدیریت محترم کلینیک،", size=13, bold=True, color=INK, space_after=8)
    add_text(
        doc,
        "در فضای پررقابت خدمات زیبایی و پزشکی، دیده شدن به تنهایی کافی نیست؛ "
        "دیده شدن توسط مخاطب هدف و آماده‌ی تصمیم‌گیری کلید اصلی موفقیت است. "
        "پلتفرم جامع «سلام دکتر» با سرمایه‌گذاری گسترده روی زیرساخت‌های سرچ و کمپین‌های قدرتمند گوگل ادز، "
        "بیمارانی را مستقیماً به سمت شما هدایت می‌کند که در لحظه، جستجوگر خدمات تخصصی شما هستند.",
    )
    add_text(doc, "پکیج طلایی (Gold VIP Plan)", size=16, bold=True, color=NAVY, space_before=6, space_after=6)
    add_text(
        doc,
        "ما برای کلینیک‌های پیشرو که قصد دارند بالاترین سهم از بازار آنلاین تخصصی خود را تصاحب کنند، "
        "پکیج طلایی را با ظرفیت بسیار محدود (تنها ۳ کلینیک در هر تخصص) طراحی کرده‌ایم.",
    )
    add_box(
        doc,
        [
            "حضور در صدر نتایج: قرارگیری کلینیک شما در ۳ جایگاه اول لندینگ‌پیج‌های تخصصی پلتفرم (پوست، مو، زیبایی و…) که بالاترین نرخ کلیک را دارند.",
            "لندینگ پیج اختصاصی: طراحی یک صفحه وب کاملاً اختصاصی و حرفه‌ای برای کلینیک شما درون پلتفرم، شامل معرفی کامل، رزومه پزشکان، اطلاعات تماس مستقیم و نقشه.",
            "تضمین ترافیک هدفمند: هدایت حداقل ۵٬۰۰۰ بازدیدکننده یونیک و جستجوگر در ماه به صفحه اختصاصی شما.",
            "برندینگ ویژه: متمایز شدن نام کلینیک با نشان VIP در لیست اصلی پلتفرم.",
        ],
        fill="FFF6D8",
        title="مزایای انحصاری پکیج طلایی (جایگاه ۱ تا ۳)",
        title_color=ORANGE,
    )
    add_text(doc, "تحلیل بازگشت سرمایه (ROI)", size=16, bold=True, color=NAVY, space_before=8, space_after=6)
    add_text(doc, "چرا این یک سرمایه‌گذاری بدون ریسک است؟", size=13, bold=True, color=INK, space_after=6)
    add_text(
        doc,
        "ارزش سرمایه‌گذاری سالانه در پکیج طلایی ۲۵۰ میلیون تومان است "
        "(با احتساب تخفیف پرداخت سالانه معادل ۲ ماه رایگان). "
        "این یعنی ماهانه حدود ۲۰٫۸ میلیون تومان برای کلینیک هزینه دارد.",
    )
    add_box(
        doc,
        [("هزینه به ازای هر بازدیدکننده هدفمند: تنها ۴٬۱۶۰ تومان", True)],
        fill="E8F4FC",
        title=None,
        body_color=BLUE,
        size=13,
    )
    add_text(
        doc,
        "منطق سودآوری: اگر از ۵٬۰۰۰ نفری که هر ماه مشخصاً دنبال خدمات زیبایی هستند و وارد صفحه شما می‌شوند، "
        "تنها ۱۵ الی ۲۵ نفر به بیمار واقعی تبدیل شوند (یک جلسه لیزر کامل، بوتاکس، یا کاشت مو)، "
        "کل هزینه سرمایه‌گذاری ماهانه شما در همان روزهای ابتدایی ماه به کلینیک بازمی‌گردد و "
        "مابقی مراجعین سود خالص مجموعه خواهند بود.",
    )
    add_box(
        doc,
        [
            "به منظور ارائه بهترین تصویر از مجموعه شما، تیم تولید محتوای پلتفرم مستقیماً در کلینیک شما حضور خواهند یافت. "
            "این تیم صرفاً وظیفه تولید محتوای بصری (عکاسی و تصویربرداری حرفه‌ای از محیط و تجهیزات) را بر عهده دارد "
            "تا گالری تصاویر لندینگ‌پیج اختصاصی شما را با بالاترین استانداردهای بصری آماده کند. "
            "این خدمات برای خریداران پکیج طلایی در این ماه کاملاً رایگان است.",
            "لازم به ذکر است فعالیت‌های مربوط به بهینه‌سازی موتورهای جستجو و زیرساخت توسط تیم فنی مجزا در بک‌اند پلتفرم مدیریت می‌شود.",
            ("+ به همراه یک ماه نمایش رایگان بنر کلینیک در صفحه اصلی پلتفرم.", True),
        ],
        fill="041F45",
        title="هدیه ویژه اعضای بنیان‌گذار (محدود به ماه اول)",
        title_color=GOLD,
        body_color=WHITE,
        size=11,
    )
    add_text(
        doc,
        "برای تنظیم جلسه حضوری و بررسی تخصصی ظرفیت‌های همکاری، با دپارتمان امور کلینیک‌های «سلام دکتر» در تماس باشید.",
        size=11,
        color=MUTED,
        align="center",
        space_before=8,
    )
    add_text(doc, "www.salam-doctor.ir", size=12, bold=True, color=BLUE, align="center")
    font_note(doc)
    path = OUT / "salam_doctor_vip_gold_pitch.docx"
    doc.save(path)
    # same file under the name first requested for the B2B pitch
    alias = OUT / "salam_doctor_b2b_pitch.docx"
    doc.save(alias)
    return path


# ---------------------------------------------------------------------------
# WA0011 — internal financial report
# ---------------------------------------------------------------------------
def build_financial():
    doc = new_doc()
    add_text(doc, "پلتفرم نوبت‌دهی و پزشکی", size=11, color=MUTED, align="center", space_after=2)
    add_text(doc, "سلام دکتر", size=22, bold=True, color=NAVY, align="center", space_after=6)
    add_text(
        doc,
        "گزارش مالی استراتژیک و لیست قیمت‌گذاری محصولات تبلیغاتی — افق زمانی ۲۴ ماهه با هدف سود خالص ۳۰٪",
        size=12,
        color=INK,
        align="center",
        space_after=6,
    )
    add_text(
        doc,
        "تاریخ گزارش: 2026/08/12  |  واحد پولی: تومان  |  نرخ دلار مبنا: 200,000 تومان",
        size=10,
        color=MUTED,
        align="center",
        space_after=10,
    )
    add_text(
        doc,
        "این گزارش بر اساس سرمایه اولیه راه‌اندازی (CAPEX)، هزینه‌های جاری ماهانه (OPEX) "
        "و هدف سودآوری ۳۰٪ روی مجموع کل هزینه‌های دو ساله تهیه شده است.",
        size=11,
    )
    add_text(doc, "۱. خلاصه اجرایی مالی", size=14, bold=True, color=NAVY, space_before=4)
    add_table(
        doc,
        ["عنوان", "مبلغ (تومان)"],
        [
            ["سرمایه اولیه (CAPEX)", "1,500,000,000"],
            ["هزینه جاری ماهانه (OPEX)", "358,000,000"],
            ["مجموع OPEX در ۲۴ ماه", "8,592,000,000"],
            ["کل سرمایه‌گذاری (CAPEX + OPEX ۲۴ماهه)", "10,092,000,000"],
            ["هدف درآمدی برای سود خالص ۳۰٪ روی هزینه دو ساله", "13,119,600,000"],
        ],
    )
    add_text(doc, "۲. جزئیات محاسبات", size=14, bold=True, color=NAVY)
    add_text(doc, "OPEX ماهانه = 1,790 دلار × 200,000 = 358,000,000 تومان", size=11, color=BLUE)
    add_table(
        doc,
        ["شرح", "فرمول / مبنا", "مبلغ (تومان)"],
        [
            ["سرمایه اولیه راه‌اندازی", "CAPEX", "1,500,000,000"],
            ["هزینه جاری ماهانه", "1,790 × 200,000", "358,000,000"],
            ["مجموع هزینه‌های جاری ۲۴ ماه", "358,000,000 × 24", "8,592,000,000"],
            ["کل سرمایه‌گذاری", "CAPEX + OPEX ۲۴ماهه", "10,092,000,000"],
            ["سود خالص هدف (۳۰٪)", "کل سرمایه‌گذاری × ۳۰٪", "3,027,600,000"],
            ["هدف درآمدی کل", "سرمایه‌گذاری + سود هدف", "13,119,600,000"],
        ],
    )
    add_text(doc, "۳. ظرفیت درآمدی یک لندینگ پیج", size=14, bold=True, color=NAVY)
    add_text(
        doc,
        "در صورت پر شدن هر ۱۰ جایگاه یک لندینگ تخصصی، ظرفیت درآمدی ماهانه "
        "(معادل‌سازی پکیج‌های یک‌ساله) به شرح زیر است:",
        size=11,
    )
    add_table(
        doc,
        ["جایگاه", "پکیج مبنا", "تعداد", "قیمت پکیج", "معادل ماهانه", "جمع ماهانه"],
        [
            ["۱ تا ۳", "طلایی ۱ ساله", "3", "250,000,000", "20,833,333", "62,500,000"],
            ["۴ تا ۶", "رشد ۱ ساله", "3", "150,000,000", "12,500,000", "37,500,000"],
            ["۷ تا ۱۰", "اقتصادی ۱ ساله", "4", "100,000,000", "8,333,333", "33,333,333"],
            ["ظرفیت (۱۰ جایگاه پر)", "ماهانه یک لندینگ", "10", "—", "—", "133,333,333"],
        ],
    )
    add_table(
        doc,
        ["شاخص", "مبلغ / مقدار"],
        [
            ["بنرهای لندینگ (۶ جایگاه)", "90,000,000 تومان"],
            ["سناریوی پکیج‌های ۶ماهه (جایگاه ۴ تا ۱۰) + طلایی سالانه", "147,500,000 تومان"],
            ["درآمد ماهانه هدف (میانگین ۲۴ ماه)", "546,650,000 تومان"],
            ["تعداد لندینگ پر لازم در ماه (تقریبی)", "حدود 4.1 لندینگ"],
        ],
        header_fill="1E4A7A",
    )
    add_text(doc, "۴. لیست قیمت‌گذاری محصولات", size=14, bold=True, color=NAVY)
    add_text(doc, "۴-۱. بنرهای تبلیغاتی ماهانه", size=12, bold=True, color=BLUE)
    add_table(
        doc,
        ["محصول", "تعداد جایگاه", "قیمت هر جایگاه", "دوره", "ظرفیت ماهانه"],
        [
            ["بنر صفحه اصلی", "3", "30,000,000", "ماهانه", "90,000,000"],
            ["بنر لندینگ پیج", "6", "15,000,000", "ماهانه", "90,000,000"],
            ["جمع — ظرفیت بنرها (فروش کامل)", "9", "—", "—", "180,000,000"],
        ],
    )
    add_text(doc, "۴-۲. پکیج‌های جایگاه در لندینگ تخصصی", size=12, bold=True, color=BLUE)
    add_table(
        doc,
        ["دسته جایگاه", "نوع پکیج", "مدت", "قیمت (تومان)", "معادل ماهانه"],
        [
            ["جایگاه ۱ تا ۳", "پکیج طلایی + صفحه اختصاصی", "12 ماه", "250,000,000", "20,833,333"],
            ["جایگاه ۴ تا ۶", "پکیج رشد", "6 ماه", "90,000,000", "15,000,000"],
            ["جایگاه ۴ تا ۶", "پکیج رشد", "12 ماه", "150,000,000", "12,500,000"],
            ["جایگاه ۷ تا ۱۰", "پکیج اقتصادی", "6 ماه", "60,000,000", "10,000,000"],
            ["جایگاه ۷ تا ۱۰", "پکیج اقتصادی", "12 ماه", "100,000,000", "8,333,333"],
        ],
    )
    add_text(doc, "۵. جمع‌بندی راهبردی", size=14, bold=True, color=NAVY)
    add_box(
        doc,
        [
            "یک لندینگ کاملاً پر، حدود 4.1 لندینگ پر در هر ماه (یا ترکیب معادل از بنر + پکیج) "
            "مسیر رسیدن به هدف ۳۰٪ سود را هموار می‌کند. در افق ۲۴ ماهه باید ترکیبی از فروش بنر و "
            "پکیج‌های جایگاه برنامه‌ریزی شود. با ظرفیت حدود 133,333,333 تومان در ماه برای هر لندینگ، "
            "رسیدن به هدف درآمدی 13,119,600,000 تومان ممکن است.",
        ],
        fill="E8F1FA",
        title=None,
        body_color=NAVY,
        size=11,
    )
    add_text(
        doc,
        "این سند صرفاً برای برنامه‌ریزی داخلی مالی و قیمت‌گذاری «سلام دکتر» تهیه شده است. "
        "فایل مبدأ: salam_doctor_financial_report.pdf / DOC-20260818-WA0011.pdf",
        size=9,
        color=MUTED,
        align="center",
        space_before=12,
    )
    font_note(doc)
    path = OUT / "salam_doctor_financial_report.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# WA0012 — small clinics / private practices
# ---------------------------------------------------------------------------
def build_small():
    doc = new_doc()
    add_text(doc, "پیشنهاد همکاری و توسعه مطب", size=20, bold=True, color=NAVY, align="center", space_after=4)
    add_text(
        doc,
        "راهکار هوشمند جذب بیمار ویژه مطب‌های شخصی و کلینیک‌های تخصصی",
        size=12,
        color=BLUE,
        align="center",
        space_after=14,
    )
    add_text(doc, "پزشک محترم، مدیریت کلینیک؛", size=13, bold=True, space_after=8)
    add_text(
        doc,
        "جذب مستمر بیمار جدید در پلتفرم‌های عمومی هزینه‌بر و اغلب بدون بازدهی دقیق است. "
        "پلتفرم «سلام دکتر» با تمرکز انحصاری بر جستجوهای مرتبط با خدمات پزشکی و زیبایی (از طریق گوگل ادز) "
        "بیماران را دقیقاً در لحظه‌ای که به خدمات شما نیاز دارند، به سمت مطب شما هدایت می‌کند. "
        "نیازی به استخدام تیم مارکتینگ ندارید؛ ما جریان ورودی بیماران را تضمین می‌کنیم.",
    )
    add_text(doc, "پکیج‌های پیشنهادی متناسب با ظرفیت شما", size=15, bold=True, color=NAVY, space_before=6)
    add_text(
        doc,
        "برای مطب‌های شخصی و کلینیک‌های در حال توسعه، دو پکیج کاملاً اقتصادی و کم‌ریسک طراحی شده است:",
        size=11,
    )
    add_text(doc, "پکیج رشد (جایگاه ۴ تا ۶)", size=14, bold=True, color=BLUE, space_before=4, space_after=4)
    add_box(
        doc,
        [
            ("سرمایه‌گذاری: ۹۰ میلیون تومان (دوره ۶ ماهه)", True),
            "ارزش ماهانه: تنها ۱۵ میلیون تومان.",
            "دیده شدن بالا: حضور در نیمه بالای لیست تخصصی شما در پلتفرم.",
            "هدیه ویژه: ۱۵ روز نمایش رایگان بنر تبلیغاتی شما در لندینگ‌پیج تخصصی.",
            "مناسب برای مطب‌هایی که ظرفیت پذیرش متوسط رو به بالا دارند و به دنبال یک جریان ثابت بیمار هستند.",
        ],
        fill="EAF6FF",
        title=None,
        body_color=INK,
    )
    add_text(doc, "پکیج ورود به بازار (جایگاه ۷ تا ۱۰)", size=14, bold=True, color=BLUE, space_after=4)
    add_box(
        doc,
        [
            ("سرمایه‌گذاری: ۶۰ میلیون تومان (دوره ۶ ماهه)", True),
            "ارزش ماهانه: تنها ۱۰ میلیون تومان.",
            "تضمین ورودی: هدایت ترافیک هدفمند با کمترین ریسک مالی ممکن.",
            "شرایط پرداخت منعطف: امکان پرداخت به‌صورت دو مرحله‌ای (۵۰٪ نقد و ۵۰٪ چک دو ماهه) ویژه پزشکانی که به‌تازگی فعالیت مستقل خود را آغاز کرده‌اند.",
            "امکان ارتقاء به پکیج‌های بالاتر در دوره‌های بعدی با ارائه تخفیف وفاداری (۱۰٪).",
        ],
        fill="EAF8F0",
        title=None,
        body_color=INK,
    )
    add_text(doc, "منطق اقتصادی این سرمایه‌گذاری برای شما", size=15, bold=True, color=NAVY, space_before=4)
    add_text(
        doc,
        "با خرید پکیج «ورود به بازار»، شما ماهانه ۱۰ میلیون تومان پرداخت می‌کنید. "
        "با احتساب هدایت حداقل ۵٬۰۰۰ بازدیدکننده جستجوگر به صفحه تخصص شما:",
        size=11,
    )
    add_box(
        doc,
        [("هزینه هر بازدیدکننده فقط ۲٬۰۰۰ تومان خواهد بود.", True)],
        fill="FDE8EF",
        body_color=RGBColor(0x9B, 0x1B, 0x4A),
        size=13,
    )
    add_text(
        doc,
        "در صنعت پزشکی، اگر از این ۵٬۰۰۰ نفر، تنها ۵ تا ۱۰ نفر به بیمار واقعی شما تبدیل شوند، "
        "کل هزینه تبلیغات ماهانه شما بازگشته است. این یعنی یک مارکتینگ کاملاً بدون باخت.",
    )
    add_text(
        doc,
        "جهت دریافت مشاوره رایگان و رزرو جایگاه با کارشناسان ما تماس بگیرید.",
        size=11,
        align="center",
        space_before=10,
    )
    footer_ir(doc)
    path = OUT / "salam_doctor_small_clinics_pitch.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# WA0013 — comprehensive financial estimate + pricing
# ---------------------------------------------------------------------------
def build_pricing():
    doc = new_doc()
    add_text(doc, "سند جامع برآورد مالی و استراتژی قیمت‌گذاری", size=18, bold=True, color=NAVY, align="center", space_after=4)
    add_text(
        doc,
        "پلتفرم نوبت‌دهی و خدمات پزشکی «سلام دکتر»",
        size=13,
        color=BLUE,
        align="center",
        space_after=12,
    )
    add_text(doc, "۱. خلاصه هزینه‌های راه‌اندازی و عملیاتی", size=14, bold=True, color=NAVY)
    add_text(
        doc,
        "با احتساب نرخ دلار ۲۰۰٬۰۰۰ تومانی، هزینه‌های پلتفرم به شرح زیر است:",
        size=11,
    )
    add_text(
        doc,
        "هزینه راه‌اندازی اولیه (CAPEX): ۱٬۵۰۰٬۰۰۰٬۰۰۰ تومان "
        "(خرید دامنه، سرور، زیرساخت، حقوق ۴ ماه تیم، گوگل ادز و تجهیزات).",
    )
    add_text(
        doc,
        "هزینه‌های جاری ماهانه (OPEX): ۱٬۷۹۰ دلار معادل ۳۵۸٬۰۰۰٬۰۰۰ تومان.",
    )
    add_text(doc, "جزئیات هزینه‌های ماهانه (دلار):", size=12, bold=True, color=INK, space_before=4)
    add_table(
        doc,
        ["آیتم", "مبلغ ماهانه (دلار)"],
        [
            ["سرور", "30"],
            ["دامنه", "10"],
            ["حقوق تیم", "1,000"],
            ["گوگل ادز", "500"],
            ["سئو", "250"],
            ["جمع", "1,790"],
        ],
    )
    add_text(doc, "۲. تعرفه تبلیغات و بنرها", size=14, bold=True, color=NAVY)
    add_table(
        doc,
        ["جایگاه", "تعداد", "قیمت ماهانه (تومان)", "توضیح"],
        [
            ["بنر صفحه اصلی", "3", "30,000,000", "بالاترین میزان دیده شدن، مناسب برای برندینگ کلینیک‌های بزرگ"],
            ["بنر لندینگ پیج", "6", "15,000,000", "نمایش هدفمند در صفحات تخصصی (پوست، مو، لاغری و…)"],
        ],
    )
    add_text(doc, "۳. استراتژی قیمت‌گذاری پلکانی جایگاه‌ها (هر لندینگ پیج)", size=14, bold=True, color=NAVY)
    add_text(
        doc,
        "هر صفحه دارای ۱۰ جایگاه است که بر اساس رتبه قیمت‌گذاری شده‌اند "
        "(با تضمین حداقل ۵٬۰۰۰ بازدید در ماه):",
        size=11,
    )
    add_table(
        doc,
        ["رتبه در صفحه", "ارزش ماهانه (تومان)", "پکیج ۶ماهه (تومان)", "پکیج ۱ساله (تومان) — پرداخت ۱۰ ماه"],
        [
            ["جایگاه ۱ تا ۳ (ویژه)", "25,000,000", "150,000,000", "250,000,000 (پلن طلایی)"],
            ["جایگاه ۴ تا ۶", "15,000,000", "90,000,000", "150,000,000"],
            ["جایگاه ۷ تا ۱۰", "10,000,000", "60,000,000", "100,000,000"],
        ],
    )
    add_text(doc, "۴. جزئیات پکیج طلایی (Gold Plan)", size=14, bold=True, color=NAVY)
    add_box(
        doc,
        [
            "این پکیج مختص جایگاه‌های ۱ تا ۳ بوده و برای کلینیک‌هایی که قصد تصاحب بیشترین سهم بازار را دارند طراحی شده است:",
            "ارزش پکیج: ۲۵۰٬۰۰۰٬۰۰۰ تومان سالانه (تخفیف ۲ ماهه به دلیل پرداخت سالانه).",
            "تضمین رتبه برتر: قرارگیری نام و مشخصات کلینیک در سه جایگاه اول لندینگ پیج.",
            "طراحی صفحه اختصاصی: ساخت لندینگ‌پیج اختصاصی برای کلینیک درون پلتفرم (شامل گالری تصاویر، خدمات و تماس مستقیم).",
            "ظرفیت تیم فنی: امکان طراحی ۴ صفحه اختصاصی در ماه (هر صفحه ۱ هفته زمان می‌برد).",
        ],
        fill="F4F7FB",
        title=None,
    )
    add_text(doc, "۵. پیش‌بینی درآمدی (ظرفیت یک لندینگ پیج)", size=14, bold=True, color=NAVY)
    add_text(
        doc,
        "در صورت تکمیل ظرفیت ۱۰۰ درصدی تنها یک لندینگ‌پیج تخصصی (بدون احتساب بنرها)، درآمد ماهانه به شرح زیر خواهد بود:",
        size=11,
    )
    add_table(
        doc,
        ["جایگاه", "تعداد", "درآمد ماهانه (تومان)"],
        [
            ["رتبه ۱ تا ۳", "3", "75,000,000"],
            ["رتبه ۴ تا ۶", "3", "45,000,000"],
            ["رتبه ۷ تا ۱۰", "4", "40,000,000"],
            ["مجموع درآمد از یک لندینگ پیج", "10", "160,000,000"],
        ],
    )
    add_text(
        doc,
        "با ۶ لندینگ‌پیج فعال، درآمد بالقوه فروش جایگاه‌ها (در حالت تکمیل ظرفیت) به ۹۶۰٬۰۰۰٬۰۰۰ تومان در ماه می‌رسد "
        "که به‌راحتی هزینه‌های جاری ۳۵۸ میلیونی را پوشش داده و سودآوری بالایی ایجاد می‌کند.",
    )
    footer_ir(doc, "فایل مبدأ: DOC-20260818-WA0013.pdf")
    path = OUT / "salam_doctor_pricing.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [build_vip(), build_financial(), build_pricing(), build_small()]
    alias = OUT / "salam_doctor_b2b_pitch.docx"
    if alias.exists() and alias not in paths:
        paths.append(alias)
    zip_path = OUT / "salam_doctor_editable_word.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        names = [
            "salam_doctor_vip_gold_pitch.docx",
            "salam_doctor_b2b_pitch.docx",
            "salam_doctor_financial_report.docx",
            "salam_doctor_pricing.docx",
            "salam_doctor_small_clinics_pitch.docx",
        ]
        for name in names:
            p = OUT / name
            if p.exists():
                zf.write(p, name)
                print(f"{name}  {p.stat().st_size}")
    print("zip", zip_path, zip_path.stat().st_size)


if __name__ == "__main__":
    main()
